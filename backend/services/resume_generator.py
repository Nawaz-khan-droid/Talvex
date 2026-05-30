"""
TALVEX Resume Generator Pipeline

End-to-end service that orchestrates the full resume generation workflow:

    1. PII Scrub    — sanitize raw resume text before any external processing
    2. Context Parse — extract structured data (contact, experience, skills, etc.)
    3. Optimize     — enhance bullets with JD keywords (when JD is provided)
    4. Template Fill — render Jinja2 HTML template with extracted data
    5. PII Rehydrate — restore real contact info into final HTML

Outputs: styled HTML, .docx bytes, Typst source, PDF bytes, ATS score,
detected skills, PII count.

Coexists with ``resume_customizer.py`` (which owns .docx template-based
customization) — this module owns the *pipeline* that strings together
templates + LLM + PII into a single generate() call.
"""

from __future__ import annotations

import io
import json
import logging
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Any, Optional

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from jinja2 import Environment

from templates.template_registry import get_template, validate_template, list_templates
from services.pii_sanitizer import sanitize_pii, restore_pii, PIISanitizer
from services.matcher import calculate_ats_score, extract_skills_from_text

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Lazy imports — these services may be built by other agents
# ---------------------------------------------------------------------------

try:
    from services.resume_optimizer import resume_optimizer as _optimizer
    _OPTIMIZER_AVAILABLE = True
except ImportError:
    _OPTIMIZER_AVAILABLE = False
    _optimizer = None
    logger.info("resume_optimizer not available — JD optimization disabled.")

try:
    from services.openrouter_client import openrouter_client as _llm_client
    _LLM_AVAILABLE = True
except ImportError:
    _LLM_AVAILABLE = False
    _llm_client = None
    logger.info("openrouter_client not available — LLM parsing disabled, using heuristics.")

# ===================================================================
# Prompt Management System — centralized prompt templates
# ===================================================================

_PROMPT_MANAGER_AVAILABLE = False

try:
    from services.prompt_manager import load_prompt, PromptNotFoundException
    _PROMPT_MANAGER_AVAILABLE = True
except ImportError:
    pass  # LLM parsing is optional; heuristic parser handles fallback


# ---------------------------------------------------------------------------
# Data types
# ---------------------------------------------------------------------------

@dataclass
class ContactInfo:
    """Structured contact information extracted from a resume."""
    name: str = ""
    phone: str = ""
    email: str = ""
    linkedin: str = ""
    portfolio: str = ""
    location: str = ""


@dataclass
class WorkExperience:
    """A single work-experience entry."""
    title: str = ""
    company: str = ""
    location: str = ""
    dates: str = ""
    bullets: list[str] = field(default_factory=list)


@dataclass
class Education:
    """A single education entry."""
    degree: str = ""
    institution: str = ""
    year: str = ""


@dataclass
class ParsedResumeData:
    """All structured data extracted from a resume."""
    contact_info: ContactInfo = field(default_factory=ContactInfo)
    professional_summary: str = ""
    work_experience: list[WorkExperience] = field(default_factory=list)
    skills: dict[str, list[str]] = field(default_factory=lambda: {"hard_skills": [], "soft_skills": []})
    education: list[Education] = field(default_factory=list)


@dataclass
class GenerationResult:
    """Final output of the resume generation pipeline."""
    html: str = ""
    docx_bytes: bytes = b""
    typst_source: str = ""
    pdf_bytes: bytes = b""
    ats_score: int = 0
    skills_detected: list[str] = field(default_factory=list)
    pii_items_count: int = 0
    template_type: str = "chronological"
    optimization_method: str = "none"

    def to_dict(self) -> dict[str, Any]:
        return {
            "html": self.html,
            "docx_bytes": self.docx_bytes,
            "typst_source": self.typst_source,
            "pdf_bytes": self.pdf_bytes,
            "ats_score": self.ats_score,
            "skills_detected": self.skills_detected,
            "pii_items_count": self.pii_items_count,
            "template_type": self.template_type,
            "optimization_method": self.optimization_method,
        }


# ---------------------------------------------------------------------------
# Heuristic resume parser
# ---------------------------------------------------------------------------

# Regex patterns for heuristic section detection
_RE_EMAIL = re.compile(
    r"[a-zA-Z0-9!#$%&'*+/=?^_`{|}~-]+(?:\.[a-zA-Z0-9!#$%&'*+/=?^_`{|}~-]+)*"
    r"@(?:[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?\.)+[a-zA-Z]{2,}"
)
_RE_PHONE = re.compile(
    r"(?:\+?\d{1,4}[\s.-]?)?(?:\(?\d{1,4}\)?[\s.-]?)?\d{1,4}[\s.-]?\d{1,4}[\s.-]?\d{1,9}"
)
_RE_LINKEDIN = re.compile(
    r"https?://(?:www\.)?linkedin\.com/in/[a-zA-Z0-9\-_/]+"
)
_RE_PORTFOLIO = re.compile(
    r"https?://(?:www\.)?[a-zA-Z0-9\-_.]+\.[a-zA-Z]{2,}(?:/[^\s]*)?"
)

_RE_SUMMARY = re.compile(
    r"(?:professional\s+summary|summary|objective|profile|about\s+me)\s*[:\-–]?\s*$",
    re.IGNORECASE,
)
_RE_EXPERIENCE = re.compile(
    r"(?:work\s+experience|professional\s+experience|employment\s+history"
    r"|work\s+history)\s*[:\-–]?\s*$",
    re.IGNORECASE,
)
_RE_SKILLS = re.compile(
    r"(?:technical\s+skills|core\s+competencies|areas\s+of\s+expertise"
    r"|technologies|proficiencies)\s*[:\-–]?\s*$",
    re.IGNORECASE,
)
_RE_EDUCATION = re.compile(
    r"(?:education|academic|qualifications|degrees)\s*[:\-–]?\s*$",
    re.IGNORECASE,
)

# Job entry heuristic: "Title at Company" or "Title – Company" patterns
# Excludes comma-only separators to avoid matching location lines like "San Francisco, CA"
_RE_JOB_ENTRY = re.compile(
    r"^\s*([A-Z][A-Za-z\s&\-']+?)\s*(?:at|—|–|-)\s*([A-Z][A-Za-z\s&\-'.]+?)\s*$"
)

# Date range heuristic: "Jan 2020 – Present" or "2020-2023"
_RE_DATE_RANGE = re.compile(
    r"(?:((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4})"
    r"\s*[-–—]\s*(Present|\d{4}|(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4})"
    r"|(\d{4})\s*[-–—]\s*(Present|\d{4}))"
)

# Education entry heuristic
_RE_EDU_ENTRY = re.compile(
    r"(?:B\.?S\.?|B\.?A\.?|M\.?S\.?|M\.?A\.?|Ph\.?D\.?|MBA|Bachelor|Master|Doctorate|Associate)"
    r"\s*(?:of|in)?\s*[\w\s&\-']+"
    r"(?:,?\s+in\s+[\w\s&\-']+)?",
    re.IGNORECASE,
)

# Bullet points
_RE_BULLET = re.compile(r"^\s*[-•*]\s*(.+)$", re.MULTILINE)

# Degree-year pattern inside education entries
_RE_DEGREE_YEAR = re.compile(r"(\d{4})")

# Common soft skills for classification
_SOFT_SKILL_KEYWORDS: set[str] = {
    "leadership", "communication", "teamwork", "collaboration", "problem-solving",
    "critical thinking", "adaptability", "time management", "project management",
    "interpersonal", "presentation", "negotiation", "mentoring", "coaching",
    "conflict resolution", "decision making", "strategic thinking", "creativity",
    "emotional intelligence", "organizational", "analytical", "detail-oriented",
}


def _classify_skill(skill: str) -> str:
    """Classify a skill as 'hard_skills' or 'soft_skills'."""
    lower = skill.lower().strip()
    if any(kw in lower for kw in _SOFT_SKILL_KEYWORDS):
        return "soft_skills"
    return "hard_skills"


def _heuristic_parse(text: str) -> ParsedResumeData:
    """
    Parse raw resume text into structured data using regex heuristics.
    Used as fallback when LLM parsing is unavailable.
    """
    data = ParsedResumeData()

    lines = text.split("\n")
    non_empty_lines = [l.strip() for l in lines if l.strip()]

    # --- Contact info (usually at the top) ---
    contact_block = "\n".join(non_empty_lines[:10])

    email_match = _RE_EMAIL.search(contact_block)
    if email_match:
        data.contact_info.email = email_match.group()

    phone_match = _RE_PHONE.search(contact_block)
    if phone_match:
        data.contact_info.phone = phone_match.group().strip()

    linkedin_match = _RE_LINKEDIN.search(contact_block)
    if linkedin_match:
        data.contact_info.linkedin = linkedin_match.group().strip()

    # Portfolio: URLs that aren't LinkedIn
    portfolio_match = _RE_PORTFOLIO.search(contact_block)
    if portfolio_match and "linkedin" not in portfolio_match.group().lower():
        data.contact_info.portfolio = portfolio_match.group().strip()

    # Name: first non-empty line that isn't an email, phone, or URL
    for line in non_empty_lines[:5]:
        if not _RE_EMAIL.match(line) and not _RE_PHONE.match(line) and not line.startswith("http"):
            # Likely a name if it's 2-4 words and doesn't contain special chars
            words = line.split()
            if 2 <= len(words) <= 5 and all(w[0].isupper() if w else False for w in words):
                data.contact_info.name = line
                break

    # --- Skills section heading (standalone "Skills" handled separately) ---
    # The standalone word "Skills" is too common in running text, so we only
    # match it as a heading when it appears on a line by itself.
    _RE_SKILLS_SOLO = re.compile(r"^skills\s*[:\-–]?\s*$", re.IGNORECASE)

    # --- Section detection ---
    section_indices: dict[str, int] = {}
    section_order = [
        ("summary", _RE_SUMMARY),
        ("experience", _RE_EXPERIENCE),
        ("skills", _RE_SKILLS),
        ("education", _RE_EDUCATION),
    ]

    for i, line in enumerate(non_empty_lines):
        for sec_name, pattern in section_order:
            if sec_name == "skills" and not pattern.match(line):
                # Fallback: try standalone "Skills" heading
                if _RE_SKILLS_SOLO.match(line):
                    section_indices.setdefault(sec_name, i)
                    break
                continue
            if pattern.match(line) and sec_name not in section_indices:
                section_indices[sec_name] = i
                break

    # Determine boundaries between sections
    sorted_sections = sorted(section_indices.items(), key=lambda x: x[1])

    def _section_lines(sec_name: str) -> list[str]:
        start = section_indices.get(sec_name, 0) + 1  # skip heading
        end_idx = None
        for name, idx in sorted_sections:
            if name == sec_name:
                continue
            if idx > section_indices.get(sec_name, 0):
                if end_idx is None or idx < end_idx:
                    end_idx = idx
        return non_empty_lines[start:end_idx]

    # --- Summary ---
    if "summary" in section_indices:
        summary_lines = _section_lines("summary")
        data.professional_summary = " ".join(summary_lines).strip()
    else:
        # Fallback: first paragraph-like block at the top
        for i, line in enumerate(non_empty_lines[:15]):
            if len(line) > 50 and not _RE_EMAIL.match(line) and not _RE_PHONE.match(line) and not line.startswith("http"):
                # Collect consecutive long lines as summary
                summary_parts = [line]
                for j in range(i + 1, min(i + 4, len(non_empty_lines))):
                    if len(non_empty_lines[j]) > 30:
                        summary_parts.append(non_empty_lines[j])
                    else:
                        break
                data.professional_summary = " ".join(summary_parts).strip()
                break

    # --- Work Experience ---
    if "experience" in section_indices:
        exp_lines = _section_lines("experience")
        current_job: WorkExperience | None = None

        for line in exp_lines:
            # Try to match job entry pattern
            job_match = _RE_JOB_ENTRY.match(line)
            date_match = _RE_DATE_RANGE.search(line)

            if job_match:
                # Save previous job
                if current_job and (current_job.title or current_job.bullets):
                    data.work_experience.append(current_job)

                current_job = WorkExperience(
                    title=job_match.group(1).strip(),
                    company=job_match.group(2).strip(),
                )
                # Check for dates on the same line
                if date_match:
                    current_job.dates = line[date_match.start():date_match.end()]
            elif date_match and current_job:
                current_job.dates = line[date_match.start():date_match.end()]
                # Remaining text might be location
                remaining = line[:date_match.start()].strip().strip(" ,")
                if remaining:
                    current_job.location = remaining
            elif line.startswith(("  ", "\t")) or _RE_BULLET.match(line):
                # Bullet point
                bullet = _RE_BULLET.sub(r"\1", line).strip()
                if bullet and current_job:
                    current_job.bullets.append(bullet)
            elif current_job and not current_job.title and len(line) > 3:
                # First line under a section heading — might be the title
                current_job.title = line.strip()

        # Don't forget the last job
        if current_job and (current_job.title or current_job.bullets):
            data.work_experience.append(current_job)

    # --- Skills ---
    if "skills" in section_indices:
        skills_lines = _section_lines("skills")
        raw_skills: list[str] = []
        for line in skills_lines:
            # Split by comma, pipe, bullet, or semicolon
            parts = re.split(r"[,;|•\-]", line)
            for part in parts:
                part = part.strip()
                if 2 <= len(part) <= 40 and not re.search(r"\d{4}", part):
                    raw_skills.append(part)

        # Classify skills
        for skill in raw_skills:
            category = _classify_skill(skill)
            data.skills[category].append(skill)

    # Also use the matcher's skill extraction on the full text
    matched_skills = extract_skills_from_text(text)
    for skill in matched_skills:
        if skill not in data.skills["hard_skills"]:
            data.skills["hard_skills"].append(skill)

    # --- Education ---
    if "education" in section_indices:
        edu_lines = _section_lines("education")
        current_edu: Education | None = None

        for line in edu_lines:
            edu_match = _RE_EDU_ENTRY.search(line)
            year_match = _RE_DEGREE_YEAR.search(line)

            if edu_match:
                if current_edu and current_edu.degree:
                    data.education.append(current_edu)
                degree_text = edu_match.group().strip()
                current_edu = Education(degree=degree_text)
                if year_match:
                    current_edu.year = year_match.group(1)
                # Check for "Degree — Institution" on same line
                after_degree = line[edu_match.end():].strip()
                inst_match = re.match(r"^[\s]*[—–-]\s*(.+)$", after_degree)
                if inst_match:
                    inst_text = inst_match.group(1).strip()
                    # Remove trailing year if present
                    inst_text = re.sub(r"\s*\d{4}\s*$", "", inst_text).strip()
                    if inst_text and not re.fullmatch(r"\d{4}", inst_text):
                        current_edu.institution = inst_text
            elif year_match and current_edu:
                current_edu.year = year_match.group(1)
            elif current_edu and not current_edu.institution and len(line) > 3:
                # Likely institution name
                if not _RE_DEGREE_YEAR.fullmatch(line.strip()):
                    current_edu.institution = line.strip()

        if current_edu and current_edu.degree:
            data.education.append(current_edu)

    return data


# ---------------------------------------------------------------------------
# LLM resume parser
# ---------------------------------------------------------------------------

_LLM_PARSE_SYSTEM_PROMPT = """\
You are an expert resume parser. Extract structured data from the resume text below.
Output ONLY valid JSON — no markdown fences, no extra text.

Return a JSON object with this exact schema:
{
  "contact_info": {
    "name": "Full Name",
    "phone": "Phone number or empty string",
    "email": "Email address or empty string",
    "linkedin": "LinkedIn URL or empty string",
    "portfolio": "Portfolio/website URL or empty string",
    "location": "City, State or empty string"
  },
  "professional_summary": "A paragraph summarizing the candidate's background",
  "work_experience": [
    {
      "title": "Job Title",
      "company": "Company Name",
      "location": "City, State",
      "dates": "Jan 2020 – Present",
      "bullets": ["Achievement bullet 1", "Achievement bullet 2"]
    }
  ],
  "skills": {
    "hard_skills": ["Python", "AWS", "React"],
    "soft_skills": ["Leadership", "Communication"]
  },
  "education": [
    {
      "degree": "B.S. Computer Science",
      "institution": "University Name",
      "year": "2020"
    }
  ]
}

Rules:
- Use empty strings for missing fields, not null.
- Extract bullets as individual strings.
- If a section is not present, use an empty array or empty string.
- Preserve the original language and capitalization.
"""


async def _llm_parse(text: str) -> ParsedResumeData | None:
    """
    Use LLM to parse resume text into structured data.
    Returns None on any failure.
    """
    if not _LLM_AVAILABLE or not _llm_client:
        return None

    # Load system prompt from prompt manager (with inline fallback)
    system_prompt = _LLM_PARSE_SYSTEM_PROMPT  # default inline fallback
    if _PROMPT_MANAGER_AVAILABLE:
        try:
            system_prompt = load_prompt("resume_parser")
        except Exception:
            pass  # Use inline fallback silently

    try:
        response = await _llm_client.call(
            model_role="parser",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": f"Parse this resume:\n\n{text[:6000]}"},
            ],
            temperature=0.1,
            max_tokens=3000,
            json_mode=True,  # Resume parser MUST return valid JSON
        )

        # Clean markdown fences if present
        cleaned = response.strip()
        if cleaned.startswith("```"):
            cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
            cleaned = re.sub(r"\s*```$", "", cleaned)

        parsed = json.loads(cleaned)

        contact = parsed.get("contact_info", {})
        skills = parsed.get("skills", {})
        return ParsedResumeData(
            contact_info=ContactInfo(
                name=contact.get("name", "") or "",
                phone=contact.get("phone", "") or "",
                email=contact.get("email", "") or "",
                linkedin=contact.get("linkedin", "") or "",
                portfolio=contact.get("portfolio", "") or "",
                location=contact.get("location", "") or "",
            ),
            professional_summary=parsed.get("professional_summary", "") or "",
            work_experience=[
                WorkExperience(
                    title=entry.get("title", "") or "",
                    company=entry.get("company", "") or "",
                    location=entry.get("location", "") or "",
                    dates=entry.get("dates", "") or "",
                    bullets=entry.get("bullets", []) or [],
                )
                for entry in (parsed.get("work_experience") or [])
            ],
            skills={
                "hard_skills": skills.get("hard_skills", []) or [],
                "soft_skills": skills.get("soft_skills", []) or [],
            },
            education=[
                Education(
                    degree=entry.get("degree", "") or "",
                    institution=entry.get("institution", "") or "",
                    year=entry.get("year", "") or "",
                )
                for entry in (parsed.get("education") or [])
            ],
        )
    except Exception as exc:
        logger.warning("LLM resume parsing failed: %s — falling back to heuristics.", exc)
        return None


# ---------------------------------------------------------------------------
# HTML → plain-text helpers
# ---------------------------------------------------------------------------

class _HTMLTextExtractor(HTMLParser):
    """Strip HTML tags and extract plain text."""

    def __init__(self) -> None:
        super().__init__()
        self._parts: list[str] = []
        self._skip = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style"):
            self._skip = True

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self._skip = False
        if tag in ("p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr"):
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self._parts.append(data)

    def get_text(self) -> str:
        return "".join(self._parts).strip()


def _html_to_text(html: str) -> str:
    """Convert HTML string to plain text."""
    extractor = _HTMLTextExtractor()
    extractor.feed(html)
    return extractor.get_text()


# ---------------------------------------------------------------------------
# ResumeGenerator
# ---------------------------------------------------------------------------

class ResumeGenerator:
    """
    End-to-end resume generation pipeline for TALVEX.

    Workflow
    --------
    1. **PII Scrub**      — sanitize_pii(resume_text) → scrubbed text + mapping
    2. **Context Parse**   — LLM or heuristic extraction of structured data
    3. **Optimize**        — (optional) resume_optimizer.optimize_resume() for JD targeting
    4. **Template Fill**   — Jinja2 render of HTML template with parsed data
    5. **PII Rehydrate**   — restore real contact info into final HTML

    Usage::

        generator = ResumeGenerator()
        result = await generator.generate(
            resume_text="John Doe\n555-1234\n...",
            job_description="Senior Python Developer...",
            template_type="chronological",
            preferences={"page_count": "1", "include_photo": False},
        )
        # result.html        → str  (styled HTML resume)
        # result.docx_bytes  → bytes (.docx file)
        # result.ats_score   → int  (0-100)
    """

    def __init__(self) -> None:
        self._pii = PIISanitizer()
        self._optimizer = _optimizer if _OPTIMIZER_AVAILABLE else None
        self._llm_available = _LLM_AVAILABLE
        self._llm_client = _llm_client

        # Jinja2 environment (no auto-escaping for HTML templates)
        self._jinja_env = Environment(
            keep_trailing_newline=True,
            trim_blocks=True,
            lstrip_blocks=True,
        )

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    async def generate(
        self,
        resume_text: str,
        job_description: str | None = None,
        template_type: str = "chronological",
        preferences: dict[str, Any] | None = None,
    ) -> GenerationResult:
        """
        Run the full resume generation pipeline.

        Parameters
        ----------
        resume_text : str
            Raw resume content (plain text or lightly formatted).
        job_description : str, optional
            Target job description. When provided, the pipeline will optimize
            resume bullets for ATS compatibility.
        template_type : str
            One of ``"chronological"``, ``"functional"``, ``"combination"``,
            ``"targeted"``.
        preferences : dict, optional
            Rendering preferences:
              - ``font_style``   (str) — CSS font family (default: Calibri, already in CSS)
              - ``page_count``   (str) — ``"1"`` or ``"2"``
              - ``include_photo``(bool) — show photo if True
              - ``photo_url``    (str) — URL of the profile photo

        Returns
        -------
        GenerationResult
            Container with ``html``, ``docx_bytes``, ``ats_score``,
            ``skills_detected``, ``pii_items_count``, ``template_type``,
            ``optimization_method``.
        """
        prefs = preferences or {}
        template_type = template_type.lower().strip()

        # Validate template type
        if not validate_template(template_type):
            valid = ", ".join(
                f'"{t["type"]}"' for t in list_templates()
            )
            logger.warning(
                "Invalid template type '%s'. Defaulting to chronological. "
                "Valid types: %s",
                template_type, valid,
            )
            template_type = "chronological"

        # ================================================================
        # STEP 0 — Extract real contact info BEFORE scrubbing
        # ================================================================
        # We need the real contact details to fill the template header.
        # After scrubbing, email/phone become placeholders, so the parser
        # cannot recover them. Extract them now and re-inject later.
        original_contact = self._pii.get_contact_info(resume_text)
        real_contact = ContactInfo(
            email=original_contact.get("email") or "",
            phone=original_contact.get("phone") or "",
            linkedin=original_contact.get("linkedin") or "",
            # GitHub URLs from PII sanitizer are stored under "github"
            portfolio=original_contact.get("github") or "",
        )
        # Heuristic: name is typically the first prominent line
        for line in resume_text.strip().split("\n")[:5]:
            line = line.strip()
            if not line:
                continue
            if _RE_EMAIL.match(line) or _RE_PHONE.match(line) or line.startswith("http"):
                continue
            words = line.split()
            if 2 <= len(words) <= 5 and all(w[0].isupper() for w in words if w):
                real_contact.name = line
                break

        # ================================================================
        # STEP 1 — PII Scrub
        # ================================================================
        scrubbed_text, pii_mapping = sanitize_pii(resume_text)
        pii_items_count = len(pii_mapping)
        logger.info(
            "Step 1 (PII Scrub): scrubbed %d PII items from resume.",
            pii_items_count,
        )

        # ================================================================
        # STEP 2 — Context Parse
        # ================================================================
        parsed_data = await self._parse_resume(scrubbed_text)

        # Re-inject real contact info (parser sees placeholders, not real PII)
        if real_contact.name:
            parsed_data.contact_info.name = real_contact.name
        if real_contact.email:
            parsed_data.contact_info.email = real_contact.email
        if real_contact.phone:
            parsed_data.contact_info.phone = real_contact.phone
        if real_contact.linkedin:
            parsed_data.contact_info.linkedin = real_contact.linkedin
        if real_contact.portfolio:
            parsed_data.contact_info.portfolio = real_contact.portfolio

        logger.info(
            "Step 2 (Parse): extracted %d jobs, %d skills, %d education entries.",
            len(parsed_data.work_experience),
            len(parsed_data.skills.get("hard_skills", []))
            + len(parsed_data.skills.get("soft_skills", [])),
            len(parsed_data.education),
        )

        # ================================================================
        # STEP 3 — Optimize (if JD provided)
        # ================================================================
        optimization_method = "none"
        if job_description and job_description.strip() and self._optimizer:
            parsed_data, optimization_method = await self._optimize_for_jd(
                parsed_data, job_description
            )
            logger.info(
                "Step 3 (Optimize): applied %s optimization.",
                optimization_method,
            )
        else:
            if job_description and job_description.strip():
                logger.info(
                    "Step 3 (Optimize): skipped — resume_optimizer not available."
                )
            else:
                logger.info("Step 3 (Optimize): skipped — no job description provided.")

        # ================================================================
        # STEP 4 — Template Fill
        # ================================================================
        page_count = str(prefs.get("page_count", "1"))
        include_photo = prefs.get("include_photo", False)
        photo_url = prefs.get("photo_url", "")

        html = self.generate_html(
            parsed_data=parsed_data,
            template_type=template_type,
            page_count=page_count,
            include_photo=include_photo and bool(photo_url),
            photo_url=photo_url,
        )
        logger.info("Step 4 (Template): rendered HTML template '%s'.", template_type)

        # ================================================================
        # STEP 5 — PII Rehydrate
        # ================================================================
        final_html = restore_pii(html, pii_mapping)
        logger.info("Step 5 (Rehydrate): restored %d PII placeholders.", pii_items_count)

        # ================================================================
        # Compute outputs
        # ================================================================
        ats_score = calculate_ats_score(
            _html_to_text(final_html),
            job_description,
        )

        all_skills = (
            parsed_data.skills.get("hard_skills", [])
            + parsed_data.skills.get("soft_skills", [])
        )
        skills_detected = list(dict.fromkeys(all_skills))  # deduplicate, preserve order

        docx_bytes = self.generate_docx(parsed_data, final_html)

        # Generate Typst source and PDF (optional — graceful fallback)
        typst_source = self.generate_typst(final_html, template_type)
        pdf_bytes = self.generate_pdf(final_html, template_type)

        return GenerationResult(
            html=final_html,
            docx_bytes=docx_bytes,
            typst_source=typst_source,
            pdf_bytes=pdf_bytes,
            ats_score=ats_score,
            skills_detected=skills_detected,
            pii_items_count=pii_items_count,
            template_type=template_type,
            optimization_method=optimization_method,
        )

    # ------------------------------------------------------------------
    # Step 2 implementation
    # ------------------------------------------------------------------

    async def _parse_resume(self, text: str) -> ParsedResumeData:
        """
        Parse scrubbed resume text into structured data.
        Tries LLM parsing first, falls back to heuristics.
        """
        # Try LLM parsing
        if self._llm_available:
            llm_result = await _llm_parse(text)
            if llm_result is not None:
                logger.info("Resume parsed via LLM.")
                # Merge with heuristic skill extraction for completeness
                heuristic_skills = extract_skills_from_text(text)
                for skill in heuristic_skills:
                    if skill not in llm_result.skills["hard_skills"]:
                        llm_result.skills["hard_skills"].append(skill)
                return llm_result

        # Heuristic fallback
        logger.info("Resume parsed via heuristics (LLM unavailable).")
        return _heuristic_parse(text)

    # ------------------------------------------------------------------
    # Step 3 implementation
    # ------------------------------------------------------------------

    async def _optimize_for_jd(
        self,
        data: ParsedResumeData,
        job_description: str,
    ) -> tuple[ParsedResumeData, str]:
        """
        Optimize resume bullets using the resume_optimizer service.
        Returns (updated_data, optimization_method).
        """
        # Reconstruct resume text from parsed data for the optimizer
        resume_text = self._data_to_text(data)

        try:
            # Determine career stage from work experience
            years = self._estimate_experience_years(data.work_experience)
            career_stage = self._career_stage_from_years(years)

            result = await self._optimizer.optimize_resume(
                resume_text=resume_text,
                job_description=job_description,
                career_stage=career_stage,
            )

            optimized_text = result.get("optimized_text", "")
            method = result.get("optimization_method", "heuristic")

            if optimized_text and len(optimized_text) > 50:
                # Re-parse the optimized text to get updated structured data
                parsed = await self._parse_resume(optimized_text)

                # Preserve original contact info (it may have been scrubbed)
                parsed.contact_info = data.contact_info

                return parsed, method
        except Exception as exc:
            logger.warning("Resume optimization failed: %s", exc)

        return data, "none"

    def _data_to_text(self, data: ParsedResumeData) -> str:
        """Convert parsed resume data back to plain text (for optimizer input)."""
        parts: list[str] = []

        # Contact
        contact = data.contact_info
        if contact.name:
            parts.append(contact.name)
        contact_items = [contact.phone, contact.email, contact.linkedin, contact.portfolio]
        contact_str = " | ".join(c for c in contact_items if c)
        if contact_str:
            parts.append(contact_str)

        # Summary
        if data.professional_summary:
            parts.append(f"\nProfessional Summary\n{data.professional_summary}")

        # Experience
        if data.work_experience:
            parts.append("\nWork Experience")
            for job in data.work_experience:
                parts.append(f"\n{job.title} at {job.company}")
                if job.location:
                    parts.append(job.location)
                if job.dates:
                    parts.append(job.dates)
                for bullet in job.bullets:
                    parts.append(f"  - {bullet}")

        # Skills
        all_skills = data.skills.get("hard_skills", []) + data.skills.get("soft_skills", [])
        if all_skills:
            parts.append(f"\nSkills\n{', '.join(all_skills)}")

        # Education
        if data.education:
            parts.append("\nEducation")
            for edu in data.education:
                parts.append(f"{edu.degree} — {edu.institution}")
                if edu.year:
                    parts.append(edu.year)

        return "\n".join(parts)

    # ------------------------------------------------------------------
    # Step 4 implementation — HTML generation
    # ------------------------------------------------------------------

    def generate_html(
        self,
        parsed_data: ParsedResumeData,
        template_type: str = "chronological",
        page_count: str = "1",
        include_photo: bool = False,
        photo_url: str = "",
    ) -> str:
        """
        Render the HTML resume template with structured data.

        Parameters
        ----------
        parsed_data : ParsedResumeData
            Structured resume content.
        template_type : str
            Template key (chronological / functional / combination / targeted).
        page_count : str
            ``"1"`` or ``"2"`` — controls CSS page-break.
        include_photo : bool
            Whether to render a profile photo ``<img>`` tag.
        photo_url : str
            URL of the profile photo image.

        Returns
        -------
        str
            Full HTML document string.
        """
        # Load template HTML
        try:
            template_html = get_template(template_type)
        except (ValueError, FileNotFoundError) as exc:
            logger.error("Failed to load template '%s': %s", template_type, exc)
            # Fallback to chronological
            try:
                template_html = get_template("chronological")
            except Exception:
                return "<html><body><p>Error loading resume template.</p></body></html>"

        # Build template context
        contact = parsed_data.contact_info
        ctx: dict[str, Any] = {
            "full_name": contact.name or "Candidate Name",
            "phone": contact.phone or "",
            "email": contact.email or "",
            "location": contact.location or "",
            "linkedin": contact.linkedin or "",
            "portfolio": contact.portfolio or "",
            "professional_summary": parsed_data.professional_summary or "",
            "work_experience": self._render_work_experience(parsed_data.work_experience),
            "skills_section": self._render_skills_section(parsed_data.skills),
            "education": self._render_education(parsed_data.education),
            "page_count": page_count,
            "include_photo": include_photo,
            "photo_url": photo_url,
        }

        # Render with Jinja2
        template = self._jinja_env.from_string(template_html)
        return template.render(**ctx)

    @staticmethod
    def _render_work_experience(jobs: list[WorkExperience]) -> str:
        """Render work experience entries as HTML."""
        if not jobs:
            return "<p style='color:#999;'>No work experience listed.</p>"

        parts: list[str] = []
        for job in jobs:
            parts.append('<div class="job-entry">')
            parts.append('  <div class="job-header">')
            parts.append(f'    <div>')
            parts.append(f'      <span class="job-title">{job.title}</span>')
            if job.company:
                parts.append(f'      <span class="job-company"> — {job.company}</span>')
            parts.append(f'    </div>')
            if job.dates:
                parts.append(f'    <span class="job-date">{job.dates}</span>')
            parts.append('  </div>')
            if job.location:
                parts.append(f'  <div class="job-location">{job.location}</div>')
            if job.bullets:
                parts.append('  <ul class="achievements">')
                for bullet in job.bullets:
                    parts.append(f'    <li>{bullet}</li>')
                parts.append('  </ul>')
            parts.append('</div>')

        return "\n".join(parts)

    @staticmethod
    def _render_skills_section(skills: dict[str, list[str]]) -> str:
        """Render skills as categorized HTML subsections."""
        hard = skills.get("hard_skills", [])
        soft = skills.get("soft_skills", [])
        if not hard and not soft:
            return "<p style='color:#999;'>No skills listed.</p>"

        parts: list[str] = []

        if hard:
            parts.append('<div class="skills-subsection">')
            parts.append('  <div class="skills-subsection-title">Technical Skills</div>')
            parts.append('  <ul class="skills-list">')
            for skill in hard:
                parts.append(f'    <li>{skill}</li>')
            parts.append('  </ul>')
            parts.append('</div>')

        if soft:
            parts.append('<div class="skills-subsection">')
            parts.append('  <div class="skills-subsection-title">Soft Skills</div>')
            parts.append('  <ul class="skills-list">')
            for skill in soft:
                parts.append(f'    <li>{skill}</li>')
            parts.append('  </ul>')
            parts.append('</div>')

        return "\n".join(parts)

    @staticmethod
    def _render_education(entries: list[Education]) -> str:
        """Render education entries as HTML."""
        if not entries:
            return "<p style='color:#999;'>No education listed.</p>"

        parts: list[str] = []
        for edu in entries:
            parts.append('<div class="edu-entry">')
            parts.append('  <div class="edu-header">')
            parts.append(f'    <div>')
            parts.append(f'      <span class="edu-degree">{edu.degree}</span>')
            if edu.institution:
                parts.append(f'      <span class="edu-school"> — {edu.institution}</span>')
            parts.append(f'    </div>')
            if edu.year:
                parts.append(f'    <span class="edu-date">{edu.year}</span>')
            parts.append('  </div>')
            parts.append('</div>')

        return "\n".join(parts)

    # ------------------------------------------------------------------
    # DOCX generation
    # ------------------------------------------------------------------

    def generate_docx(
        self,
        parsed_data: ParsedResumeData,
        html_content: str,
    ) -> bytes:
        """
        Generate a .docx file from the parsed resume data.

        Uses python-docx to create a cleanly formatted document. The HTML
        content is used as a reference, but the actual document is built
        programmatically for reliable formatting.

        Parameters
        ----------
        parsed_data : ParsedResumeData
            Structured resume content.
        html_content : str
            Rendered HTML (unused for layout but kept for future html→docx
            conversion support).

        Returns
        -------
        bytes
            The .docx file contents.
        """
        doc = Document()

        # --- Default font ---
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        # --- Narrow margins ---
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        contact = parsed_data.contact_info

        # === Header: Name ===
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(contact.name or "Candidate Name")
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)

        # === Header: Contact line ===
        contact_items: list[str] = []
        if contact.location:
            contact_items.append(contact.location)
        if contact.phone:
            contact_items.append(contact.phone)
        if contact.email:
            contact_items.append(contact.email)
        if contact.linkedin:
            contact_items.append(contact.linkedin)
        if contact.portfolio:
            contact_items.append(contact.portfolio)

        if contact_items:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(" | ".join(contact_items))
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        # Divider
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("─" * 80)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)

        # === Professional Summary ===
        if parsed_data.professional_summary:
            self._add_section_heading(doc, "Professional Summary")
            doc.add_paragraph(parsed_data.professional_summary)

        # === Work Experience ===
        if parsed_data.work_experience:
            self._add_section_heading(doc, "Work Experience")
            for job in parsed_data.work_experience:
                # Title + Company
                p = doc.add_paragraph()
                run = p.add_run(job.title)
                run.bold = True
                run.font.size = Pt(12)
                if job.company:
                    run = p.add_run(f" — {job.company}")

                # Location + Dates
                meta_parts = []
                if job.location:
                    meta_parts.append(job.location)
                if job.dates:
                    meta_parts.append(job.dates)
                if meta_parts:
                    p = doc.add_paragraph()
                    run = p.add_run(" | ".join(meta_parts))
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                    run.italic = True

                # Bullets
                for bullet in job.bullets:
                    doc.add_paragraph(bullet, style="List Bullet")

                # Spacer
                doc.add_paragraph("")

        # === Skills ===
        hard = parsed_data.skills.get("hard_skills", [])
        soft = parsed_data.skills.get("soft_skills", [])
        if hard or soft:
            self._add_section_heading(doc, "Skills")

            if hard:
                p = doc.add_paragraph()
                run = p.add_run("Technical Skills: ")
                run.bold = True
                run = p.add_run(", ".join(hard))
                run.font.size = Pt(11)

            if soft:
                p = doc.add_paragraph()
                run = p.add_run("Soft Skills: ")
                run.bold = True
                run = p.add_run(", ".join(soft))
                run.font.size = Pt(11)

        # === Education ===
        if parsed_data.education:
            self._add_section_heading(doc, "Education")
            for edu in parsed_data.education:
                p = doc.add_paragraph()
                run = p.add_run(edu.degree)
                run.bold = True
                if edu.institution:
                    run = p.add_run(f" — {edu.institution}")
                if edu.year:
                    run = p.add_run(f" ({edu.year})")
                    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        # --- Write to bytes ---
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def _add_section_heading(doc: Document, text: str) -> None:
        """Add a styled section heading to the document."""
        h = doc.add_heading(text, level=1)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
            run.font.size = Pt(14)

    # ------------------------------------------------------------------
    # Typst generation
    # ------------------------------------------------------------------

    def generate_typst(
        self,
        html: str,
        template_type: str = "chronological",
    ) -> str:
        """
        Convert rendered HTML resume to Typst markup.

        Parameters
        ----------
        html : str
            Rendered HTML resume (after PII rehydration).
        template_type : str
            Template key for the Typst template to use.

        Returns
        -------
        str
            Typst source markup. Returns empty string on any error.
        """
        try:
            from services.typst_compiler import html_to_typst
            typst_src = html_to_typst(html, template_type)
            logger.info(
                "Generated Typst source (%d chars) for template '%s'.",
                len(typst_src), template_type,
            )
            return typst_src
        except Exception as exc:
            logger.warning("Typst source generation failed: %s", exc)
            return ""

    def generate_pdf(
        self,
        html: str,
        template_type: str = "chronological",
    ) -> bytes:
        """
        Generate PDF bytes from rendered HTML resume via Typst.

        If Typst CLI is not available, returns empty bytes (caller
        should fall back to browser-based print-to-PDF).

        Parameters
        ----------
        html : str
            Rendered HTML resume (after PII rehydration).
        template_type : str
            Template key for Typst template selection.

        Returns
        -------
        bytes
            PDF document bytes, or ``b""`` if Typst is unavailable.
        """
        try:
            from services.typst_compiler import html_to_pdf
            pdf = html_to_pdf(html, template_type)
            if pdf:
                logger.info("Generated PDF (%d bytes) via Typst.", len(pdf))
            else:
                logger.info("Typst PDF generation returned empty — Typst may not be available.")
            return pdf or b""
        except Exception as exc:
            logger.warning("PDF generation failed: %s", exc)
            return b""

    # ------------------------------------------------------------------
    # Career stage & template recommendation
    # ------------------------------------------------------------------

    @staticmethod
    def _estimate_experience_years(jobs: list[WorkExperience]) -> float:
        """Estimate total years of work experience from parsed jobs."""
        import re as _re
        total_years = 0.0
        current_year = 2025

        for job in jobs:
            dates = job.dates or ""
            # Try to extract year ranges
            year_matches = _re.findall(r"\b(19|20)\d{2}\b", dates)
            if len(year_matches) >= 2:
                start = int(year_matches[0])
                end_str = year_matches[1]
                end = int(end_str) if int(end_str) <= current_year else current_year
                total_years += max(0, end - start)
            elif len(year_matches) == 1:
                start = int(year_matches[0])
                total_years += max(0, current_year - start)

        return total_years

    @staticmethod
    def _career_stage_from_years(years: float) -> str:
        """Map experience years to a career stage string."""
        if years < 2:
            return "entry_level"
        elif years < 6:
            return "mid_level"
        elif years < 12:
            return "senior"
        else:
            return "executive"

    @staticmethod
    def recommend_template(
        work_experience_years: float = 0,
        career_gaps: int = 0,
        career_changes: int = 0,
        has_specific_jd: bool = False,
    ) -> str:
        """
        Recommend the best resume template based on career profile.

        Parameters
        ----------
        work_experience_years : float
            Estimated total years of professional experience.
        career_gaps : int
            Number of significant employment gaps (> 6 months).
        career_changes : int
            Number of major career/industry changes.
        has_specific_jd : bool
            Whether a specific job description is being targeted.

        Returns
        -------
        str
            One of ``"chronological"``, ``"functional"``, ``"combination"``,
            or ``"targeted"``.

        Recommendation logic
        --------------------
        - ``< 2`` years experience or career change → **functional**
        - ``2–5`` years experience → **combination**
        - ``5+`` years with steady progression → **chronological**
        - Highly specific JD provided → **targeted** (overrides others)
        - Multiple career gaps → **functional** or **combination**
        """
        # Targeted takes highest priority if JD is very specific
        if has_specific_jd:
            return "targeted"

        # Career gaps or changes suggest functional
        if career_changes > 0 or work_experience_years < 2:
            return "functional"

        if career_gaps > 1:
            return "functional"

        # 2-5 years → combination (hybrid approach)
        if work_experience_years < 5:
            return "combination"

        # 5+ years steady progression → chronological
        return "chronological"


# ---------------------------------------------------------------------------
# Module-level singleton
# ---------------------------------------------------------------------------

resume_generator = ResumeGenerator()
